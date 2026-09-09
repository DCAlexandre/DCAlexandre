# -*- coding: utf-8 -*-
"""
Génère la version ATS du CV (1 colonne, texte réel, sans tableau ni icône).

Sortie : public/assets/cv/CV-Alexandre-Da-Costa-ATS.docx

Dépendances :
    python -m pip install python-docx

Usage (depuis la racine du repo) :
    python scripts/generate-cv-ats.py

Produire le PDF à partir du DOCX (Windows + Word installé), en PowerShell :
    $d = New-Object -ComObject Word.Application
    $doc = $d.Documents.Open("$PWD\\public\\assets\\cv\\CV-Alexandre-Da-Costa-ATS.docx",$false,$true)
    $doc.ExportAsFixedFormat("$PWD\\public\\assets\\cv\\CV-Alexandre-Da-Costa-ATS.pdf",17)
    $doc.Close($false); $d.Quit()
    # (ouvrir le .docx depuis le repo, pas depuis %TEMP% — sinon lecture protégée = blocage)

La source de vérité du CONTENU est ce fichier ; il doit rester cohérent avec le
CV design public/files/cvalexandredacosta.html.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- CONTENU (source unique)
NAME = "Alexandre Da Costa"
TITLE = "Tech Lead — Architecture & industrialisation (web, mobile, desktop)"
CONTACT = [
    "Paris et périphérie · remote ou hybride",
    "07 69 62 43 79 · alexandre@kared-dev.fr",
    "Portfolio : kared-dev.fr/alexandre · LinkedIn : linkedin.com/in/alexandre-dacosta · GitHub : github.com/DCAlexandre",
]
PROFILE = ("Tech Lead, 10 ans d'expérience. Je conçois, reprends et industrialise des applications métier, "
           "de l'idée à la production — web, mobile et desktop. J'interviens là où les projets ralentissent : "
           "dette technique, montée en charge, déploiements à risque. J'encadre aussi les équipes qui livrent. "
           "10 ans d'expérience dont 6 en lead · 4 développeurs encadrés · 100 000+ utilisateurs (Comète Link) · "
           "déploiement en 15 min au lieu de 2 h.")

SKILLS = [
    ("Leadership & méthodes", "encadrement & mentorat, revue de code, cadrage produit, priorisation / roadmap, Agile / Scrum, relation client"),
    ("Architecture & DevOps", "architecture web scalable, modernisation de legacy, temps réel, monorepo, APIs REST, CI/CD, Docker, automatisation, sécurité serveurs"),
    ("Front & mobile", "TypeScript, JavaScript, React, React Native, Ionic, Capacitor, Electron, UX produit, responsive"),
    ("Back-end & données", "Node.js, NestJS, Laravel, PostgreSQL, Prisma, Socket.IO, Python, Omnis"),
    ("IA & outillage", "modèle de langage en local, RAG documentaire, agents via MCP, Git"),
]

EXPERIENCE = [
    {
        "title": "Tech Lead / Lead Web Developer",
        "meta": "Logiciel Comète — Région parisienne — 01/2020 à aujourd'hui",
        "ctx": "Éditeur de logiciels métiers pour la sécurité privée (650 clients).",
        "bullets": [
            "Fondateur technique de l'écosystème Comète Link : créateur des 10 dépôts, auteur majoritaire de 7 d'entre eux (67 à 97 % des commits) depuis 2020.",
            "Encadrement de 4 développeurs : estimations, répartition des tickets, revues de code, suivi quotidien.",
            "Industrialisation de la chaîne de livraison : CI/CD (tests, builds, déploiements, signature), environnements Dev / Pré-prod / Prod. Déploiement en 15 min au lieu de 2 h, plus aucun build manuel.",
            "Pilotage de la reprise du legacy Omnis vers une stack web / mobile unifiée et du support de production.",
            "Projets clés : Comète Link (webapp métier modulaire, 100 000+ utilisateurs, desktop / tablette / mobile) ; Orca (moteur de synchronisation temps réel, Electron / Node / React) ; Comète Cloud & DevTools (outillage interne : installation et mise à jour sans ligne de commande, builds et publication stores en un clic).",
        ],
        "stack": "React, Ionic, Node.js, NestJS, Laravel, PostgreSQL, Socket.IO, Electron, Docker, CI/CD, Omnis.",
    },
    {
        "title": "Tech Lead freelance",
        "meta": "Kared Dev — Paris — 12/2024 à aujourd'hui",
        "ctx": "Conception et reprise d'applications web & mobile sur mesure (start-ups, PME, indépendants), en parallèle du poste chez Comète.",
        "bullets": [
            "Une seule base de code web + mobile (React / React Native / Capacitor) : un budget, une équipe, des évolutions déployées partout à la fois.",
            "5 clients accompagnés depuis décembre 2024.",
            "Projets clés : Dokart (SaaS de gestion pour studios de tatouage, reconstruction from scratch, bêta publique en 2 mois, 8 h d'administratif économisées par semaine et 98 % de no-shows en moins) ; Solutions Terrains (app mobile de sourcing foncier, publiée iOS & Android en 6 semaines, 420 comptes) ; Kared Fit (app mobile de fitness communautaire, temps réel).",
        ],
        "stack": "React Native, Capacitor, React, NestJS, PostgreSQL, Prisma, Docker, CI/CD.",
    },
    {
        "title": "Software Engineer, puis Développeur logiciels",
        "meta": "AEXAE — Les Ulis (91) — 01/2017 à 01/2020",
        "ctx": "Éditeur de logiciels métiers (sécurité privée, formation professionnelle, gestion interne).",
        "bullets": [
            "Développement et maintenance de Comète, ERP complet pour la sécurité privée (RH, facturation, plannings, paie). Prise en main d'Omnis en autonomie.",
            "Contribution à Isiform (organismes de formation) et amélioration de PMH, l'outil interne de ticketing.",
            "Refonte de l'outil web de gestion des licences : application fullstack avec tableaux de bord d'usage.",
            "Guidelines UI/UX sur toute la gamme ; support technique avancé sur les serveurs clients.",
        ],
        "stack": None,
    },
    {
        "title": "Développeur web (premières expériences)",
        "meta": "Indépendant / GuestWhat — 10/2015 à 01/2017",
        "ctx": None,
        "bullets": [
            "Indépendant (2016-2017) : sites et outils sur mesure pour TPE et start-ups, de la maquette à la mise en ligne.",
            "GuestWhat (2015-2016) : deux plateformes web, administration de contenu, déploiement et SEO.",
        ],
        "stack": None,
    },
]

EDUCATION = [
    "Docker intensif — Conteneurisation & DevOps — Enix, 2025",
    "BTS SIO (Services informatiques aux organisations) — INSTA, 2015–2017",
    "Licence Informatique — Université d'Évry Val-d'Essonne, 2013–2014",
    "Certifications : Claude Code in Action, Claude Code 101, Introduction to Model Context Protocol",
]
LANGUAGES = ("Français (langue maternelle) · Portugais (courant) · "
             "Espagnol (compréhension courante, oral intermédiaire) · Anglais (compréhension courante, oral intermédiaire)")
INTERESTS = ("Serveur auto-hébergé : modèle de langage en local, RAG sur documents personnels, "
             "agents connectés via MCP, pilotage de la domotique.")

INK = RGBColor(0x1a, 0x1a, 0x1a)
EMERALD = RGBColor(0x15, 0x7A, 0x48)   # accent marque (contraste OK sur blanc)
EMERALD_HEX = "157A48"


# ---------------------------------------------------------------- DOCX
def _set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    pf = st.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = 1.08


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = EMERALD
    # filet bas (bordure de paragraphe, PAS un tableau → reste ATS-safe)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), EMERALD_HEX)
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def build_docx(path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.4)
        s.left_margin = s.right_margin = Cm(1.6)
    _set_base_style(doc)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(NAME); r.bold = True; r.font.size = Pt(21); r.font.color.rgb = EMERALD
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(TITLE); r.font.size = Pt(11.5); r.bold = True; r.font.color.rgb = INK
    for line in CONTACT:
        cp = doc.add_paragraph(); cp.paragraph_format.space_after = Pt(0)
        cr = cp.add_run(line); cr.font.size = Pt(9.5); cr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    _heading(doc, "Profil")
    doc.add_paragraph(PROFILE)

    _heading(doc, "Compétences")
    for cat, items in SKILLS:
        sp = doc.add_paragraph()
        b = sp.add_run(cat + " : "); b.bold = True
        sp.add_run(items)

    _heading(doc, "Expérience")
    for j in EXPERIENCE:
        tp = doc.add_paragraph(); tp.paragraph_format.space_before = Pt(4); tp.paragraph_format.space_after = Pt(0)
        tr = tp.add_run(j["title"]); tr.bold = True; tr.font.size = Pt(11)
        mp = doc.add_paragraph(); mp.paragraph_format.space_after = Pt(1)
        mr = mp.add_run(j["meta"]); mr.font.size = Pt(9.5); mr.italic = True
        if j["ctx"]:
            cp = doc.add_paragraph(); cp.paragraph_format.space_after = Pt(1)
            cr = cp.add_run(j["ctx"]); cr.font.size = Pt(9.5); cr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        for b in j["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
        if j["stack"]:
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
            sr = sp.add_run("Stack : "); sr.bold = True; sr.font.size = Pt(9.5)
            sp.add_run(j["stack"]).font.size = Pt(9.5)

    _heading(doc, "Formation")
    for e in EDUCATION:
        doc.add_paragraph(e, style="List Bullet")

    _heading(doc, "Langues")
    doc.add_paragraph(LANGUAGES)

    _heading(doc, "Centres d'intérêt")
    doc.add_paragraph(INTERESTS)

    doc.save(path)
    print("DOCX ->", path)


if __name__ == "__main__":
    import os
    out_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "public", "assets", "cv"))
    os.makedirs(out_dir, exist_ok=True)
    build_docx(os.path.join(out_dir, "CV-Alexandre-Da-Costa-ATS.docx"))
    print("OK. Générer le PDF depuis le .docx via Word (voir le docstring en tête de fichier).")
